"""Regression tests against a disposable schema in the configured PostgreSQL DB.

Run: python -m unittest discover -s tests -v
External providers are mocked; extraction, JWT, SQL, and audio stitching are real.
"""
import asyncio
import io
import tempfile
import unittest
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import fitz
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session, configure_mappers
from alembic import command
from alembic.config import Config

from core.config import settings
from db.session import get_db
from models import User, ConversionHistory, Document, SystemLog
from models.user import UserRole
from services.system_log_service import log_system_event
from app.main import app


class DatabaseExpansionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.schema = "test_documents_" + uuid.uuid4().hex
        cls.control = create_engine(settings.DATABASE_URL)
        with cls.control.begin() as conn:
            conn.execute(text(f'CREATE SCHEMA "{cls.schema}"'))
        cls.engine = create_engine(settings.DATABASE_URL, connect_args={"options": f"-csearch_path={cls.schema}"})
        cls.addClassCleanup(cls.cleanup_schema)
        cls.migrate("upgrade", "00ab090b8bdc")
        with cls.engine.begin() as conn:
            conn.execute(text("INSERT INTO users (id, name, email, password, role) VALUES (9000, 'Legacy', 'legacy@example.com', 'unused', 'USER')"))
            conn.execute(text("INSERT INTO conversion_history (id,user_id,title,text,language,voice,audio_url,source_type) VALUES (9000,9000,'Legacy','Preserve me','twi','Kofi','/audio/legacy.wav','text-to-speech')"))
            cls.legacy_user = dict(conn.execute(text("SELECT * FROM users WHERE id=9000")).mappings().one())
            cls.legacy_conversion = dict(conn.execute(text("SELECT * FROM conversion_history WHERE id=9000")).mappings().one())
        cls.migrate("upgrade", "head")
        configure_mappers()

        def override_db():
            with Session(cls.engine) as db:
                yield db
        app.dependency_overrides[get_db] = override_db
        cls.addClassCleanup(app.dependency_overrides.clear)
        # No lifespan: BERT and network provider startup are tested separately.
        cls.client = TestClient(app)
        cls.addClassCleanup(cls.client.close)
        cls.owner = cls.register("owner")
        cls.other = cls.register("other")
        cls.admin = cls.register("admin")
        with Session(cls.engine) as db:
            user = db.get(User, cls.admin["user"]["id"])
            user.role = UserRole.ADMIN
            db.commit()

    @classmethod
    def cleanup_schema(cls):
        cls.engine.dispose()
        with cls.control.begin() as conn:
            # Generated schema name is fixed to this test's UUID, never public.
            conn.execute(text(f'DROP SCHEMA "{cls.schema}" CASCADE'))
        cls.control.dispose()

    @classmethod
    def migrate(cls, direction, revision):
        config = Config("alembic.ini")
        with cls.engine.begin() as conn:
            config.attributes["connection"] = conn
            getattr(command, direction)(config, revision)

    @classmethod
    def register(cls, name):
        response = cls.client.post("/auth/register", json={"name": name, "email": f"{name}@example.com", "password": "test-password-123"})
        if response.status_code != 201:
            raise AssertionError(f"Registration failed: {response.status_code}")
        return response.json()

    def headers(self, user=None):
        return {"Authorization": "Bearer " + (user or self.owner)["access_token"]}

    def upload(self, name="notes.txt", data=b"Scientific document content.", mime="text/plain"):
        return self.client.post("/upload", headers=self.headers(), files={"file": (name, data, mime)})

    def save_conversion(self, document_id=None):
        return self.client.post("/library/", headers=self.headers(), json={
            "title": "Saved conversion", "text": "Scientific document content.", "voice": "Kofi",
            "language": "twi", "source_type": "text-to-speech", "audio_url": "/audio/test.wav",
            "document_id": document_id,
        })

    def test_migration_preserves_legacy_rows(self):
        with self.engine.connect() as conn:
            self.assertEqual(dict(conn.execute(text("SELECT * FROM users WHERE id=9000")).mappings().one()), self.legacy_user)
            row = dict(conn.execute(text("SELECT * FROM conversion_history WHERE id=9000")).mappings().one())
            self.assertIsNone(row.pop("document_id"))
            self.assertEqual(row, self.legacy_conversion)

    def test_auth_login_jwt_and_admin_access(self):
        response = self.client.post("/auth/login", json={"email": "owner@example.com", "password": "test-password-123"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.get("/auth/me", headers=self.headers()).status_code, 200)
        self.assertEqual(self.client.get("/documents").status_code, 401)
        self.assertEqual(self.client.get("/documents", headers={"Authorization": "Bearer invalid"}).status_code, 401)
        self.assertEqual(self.client.get("/admin/logs", headers=self.headers()).status_code, 403)
        self.assertEqual(self.client.get("/admin/stats", headers=self.headers(self.admin)).status_code, 200)

    def test_retired_config_routes_and_existing_tts_contract(self):
        self.assertEqual(self.client.get("/admin/config", headers=self.headers(self.admin)).status_code, 404)
        self.assertEqual(self.client.put("/admin/config", headers=self.headers(self.admin), json={}).status_code, 404)
        self.assertNotIn("/admin/config", app.openapi()["paths"])
        payload = {"text": "Speech input", "language": "twi", "voice": "Kofi", "speed": 1.2, "pitch": 2}
        with patch("services.tts_service.process_tts", return_value=("/audio/test.wav", 1.0)) as processor:
            response = self.client.post("/tts", headers=self.headers(), json=payload)
            self.assertEqual(response.status_code, 200)
            processor.assert_called_once_with("Speech input", "twi", "Kofi")
            self.assertEqual(set(response.json()), {"audio_url", "duration_seconds", "language", "voice", "document_id"})

    def test_upload_supported_formats_and_compatibility(self):
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Scientific PDF content.")
        pdf_bytes = pdf.tobytes()
        pdf.close()
        docx = DocxDocument()
        docx.add_paragraph("Scientific DOCX content.")
        stream = io.BytesIO()
        docx.save(stream)
        for name, data, mime in [
            ("notes.txt", b"Scientific text content.", "text/plain"),
            ("notes.pdf", pdf_bytes, "application/pdf"),
            ("notes.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ]:
            with self.subTest(name=name):
                response = self.upload(name, data, mime)
                self.assertEqual(response.status_code, 200, response.text)
                result = response.json()
                self.assertTrue({"file_name", "text", "file_size", "file_type", "word_count", "document_id", "title", "source_type", "extracted_text"} <= result.keys())
                self.assertEqual(result["text"], result["extracted_text"])
                with Session(self.engine) as db:
                    document = db.get(Document, result["document_id"])
                    self.assertEqual(document.user_id, self.owner["user"]["id"])
                    self.assertIsNone(document.file_url)
                    self.assertIsNotNone(document.created_at.tzinfo)

    def test_failed_extraction_leaves_no_document(self):
        with Session(self.engine) as db:
            before = db.query(Document).count()
        self.assertEqual(self.upload(data=b"   ").status_code, 422)
        self.assertEqual(self.upload("image.png", b"image", "image/png").status_code, 400)
        from fastapi import HTTPException
        with patch("api_routes.upload.extract_text_from_pdf", side_effect=HTTPException(422, "Extraction failed")):
            self.assertEqual(self.upload("bad.pdf", b"invalid", "application/pdf").status_code, 422)
        with Session(self.engine) as db:
            self.assertEqual(db.query(Document).count(), before)

    def test_url_scrape_persistence(self):
        article = SimpleNamespace(title="Article title", text="Readable article content.", download=lambda: None, parse=lambda: None)
        with patch("services.scrape_service.Article", return_value=article):
            response = self.client.post("/scrape/", headers=self.headers(), json={"url": "https://example.com/article"})
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        with Session(self.engine) as db:
            document = db.get(Document, result["document_id"])
            self.assertEqual(document.source_type, "url")
            self.assertEqual(document.source_url, "https://example.com/article")
            self.assertEqual(document.title, "Article title")
            self.assertEqual(document.extracted_text, result["text"])

    def test_ownership_and_safe_document_deletion(self):
        document_id = self.upload().json()["document_id"]
        response = self.client.get("/documents", headers=self.headers())
        self.assertEqual(response.status_code, 200)
        self.assertIn(document_id, [d["id"] for d in response.json()])
        self.assertNotIn("extracted_text", response.json()[0])
        self.assertEqual(self.client.get(f"/documents/{document_id}", headers=self.headers()).status_code, 200)
        for user in [self.other, self.admin]:
            self.assertEqual(self.client.get(f"/documents/{document_id}", headers=self.headers(user)).status_code, 403)
            self.assertEqual(self.client.delete(f"/documents/{document_id}", headers=self.headers(user)).status_code, 403)
            self.assertNotIn(document_id, [d["id"] for d in self.client.get("/documents", headers=self.headers(user)).json()])
        history = self.save_conversion(document_id)
        self.assertEqual(history.status_code, 201, history.text)
        self.assertEqual(history.json()["document_id"], document_id)
        self.assertEqual(self.client.delete(f"/documents/{document_id}", headers=self.headers()).status_code, 200)
        with Session(self.engine) as db:
            self.assertIsNone(db.get(Document, document_id))
            conversion = db.get(ConversionHistory, history.json()["id"])
            self.assertIsNotNone(conversion)
            self.assertIsNone(conversion.document_id)
            events = db.query(SystemLog).filter(SystemLog.user_id == self.owner["user"]["id"]).all()
            self.assertTrue(any(e.action == "DOCUMENT_UPLOADED" and f"document {document_id} " in e.details for e in events))
            self.assertTrue(any(e.action == "DOCUMENT_DELETED" and e.details == f"Deleted document {document_id}" for e in events))
        self.assertEqual(self.client.get(f"/documents/{document_id}", headers=self.headers()).status_code, 404)

    def test_tts_document_validation_and_separate_history_save(self):
        document_id = self.upload().json()["document_id"]
        payload = {"text": "Scientific input.", "language": "twi", "voice": "Kofi"}
        with patch("services.tts_service.process_tts", return_value=("/audio/test.wav", 1.0)) as processor:
            for ref in [None, document_id]:
                response = self.client.post("/tts/", headers=self.headers(), json={**payload, "document_id": ref})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["document_id"], ref)
                history = self.save_conversion(ref)
                self.assertEqual(history.status_code, 201)
                self.assertEqual(history.json()["document_id"], ref)
            self.assertEqual(processor.call_count, 2)
            response = self.client.post("/tts/", headers=self.headers(self.other), json={**payload, "document_id": document_id})
            self.assertEqual(response.status_code, 403)
            self.assertEqual(self.client.post("/tts/", headers=self.headers(), json={**payload, "document_id": 999999}).status_code, 404)
            self.assertEqual(processor.call_count, 2)
        response = self.client.post("/library/", headers=self.headers(self.other), json={**payload, "title": "Other", "source_type": "text", "document_id": document_id})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.save_conversion(999999).status_code, 404)

    def test_history_bookmark_delete(self):
        conversion = self.save_conversion().json()
        conversion_id = conversion["id"]
        self.assertIsNone(conversion["document_id"])
        self.assertEqual(self.client.get("/library/", headers=self.headers()).status_code, 200)
        self.assertEqual(self.client.post(f"/library/{conversion_id}/bookmark", headers=self.headers()).json(), {"bookmarked": True})
        self.assertEqual(self.client.post(f"/library/{conversion_id}/bookmark", headers=self.headers()).json(), {"bookmarked": False})
        self.assertEqual(self.client.delete(f"/library/{conversion_id}", headers=self.headers(self.other)).status_code, 403)
        self.assertEqual(self.client.delete(f"/library/{conversion_id}", headers=self.headers()).status_code, 200)

    def test_audit_events_admin_changes_pagination_no_access_logs(self):
        target = self.other["user"]["id"]
        for update in [{"is_banned": True}, {"is_banned": False}, {"role": "admin"}, {"role": "user"}]:
            self.assertEqual(self.client.put(f"/admin/users/{target}", headers=self.headers(self.admin), json=update).status_code, 200)
        with Session(self.engine) as db:
            events = db.query(SystemLog).filter(SystemLog.user_id == self.admin["user"]["id"]).all()
            self.assertTrue({"USER_BANNED", "USER_UNBANNED", "ROLE_CHANGED"} <= {e.action for e in events})
            before = db.query(SystemLog).count()
        response = self.client.get("/admin/logs?limit=2&offset=0", headers=self.headers(self.admin))
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual(len(result["logs"]), 2)
        self.assertGreater(result["logs"][0]["id"], result["logs"][1]["id"])
        self.assertTrue({"timestamp", "level", "message", "user_email", "action", "details", "created_at"} <= result["logs"][0].keys())
        self.assertEqual(self.client.get("/admin/logs?limit=201", headers=self.headers(self.admin)).status_code, 422)
        with Session(self.engine) as db:
            self.assertEqual(db.query(SystemLog).count(), before)
            self.assertTrue(log_system_event(db, None, "SYSTEM_TEST", "No authenticated user"))

    def test_audit_failure_does_not_rollback_document(self):
        document_id = self.upload().json()["document_id"]
        with Session(self.engine) as db:
            self.assertFalse(log_system_event(db, 999999, "TEST_FAILURE", "Invalid user reference"))
            self.assertIsNotNone(db.get(Document, document_id))
            self.assertEqual(db.execute(text("SELECT 1")).scalar(), 1)

    def test_main_activity_audit_success_and_failure(self):
        from fastapi import HTTPException

        def actions():
            with Session(self.engine) as db:
                return [row.action for row in db.query(SystemLog).filter(
                    SystemLog.user_id == self.owner["user"]["id"]).order_by(SystemLog.id).all()]

        self.assertIn("USER_REGISTERED", actions())
        before = len(actions())
        self.assertEqual(self.client.post("/auth/login", json={"email": "owner@example.com", "password": "wrong"}).status_code, 401)
        self.assertEqual(len(actions()), before)
        self.assertEqual(self.client.post("/auth/login", json={"email": "owner@example.com", "password": "test-password-123"}).status_code, 200)
        self.assertEqual(actions()[-1], "USER_LOGGED_IN")
        self.assertEqual(self.client.post("/auth/register", json={"name": "Duplicate", "email": "owner@example.com", "password": "test-password-123"}).status_code, 400)
        self.assertEqual(len(actions()), before + 1)

        conversion = self.save_conversion().json()
        self.assertEqual(actions()[-1], "CONVERSION_CREATED")
        for expected in ["CONVERSION_BOOKMARKED", "CONVERSION_UNBOOKMARKED"]:
            self.assertEqual(self.client.post(f"/library/{conversion['id']}/bookmark", headers=self.headers()).status_code, 200)
            self.assertEqual(actions()[-1], expected)

        cases = [
            ("/tts/", "services.tts_service.process_tts", {"text": "Private input", "language": "twi", "voice": "Kofi"}, ("/audio/test.wav", 1.0), "TTS_GENERATED"),
            ("/summarise", "services.ai_service.summarise_text", {"text": "Private input", "language": "twi"}, {"summary": "Private summary", "word_count": 2}, "SUMMARY_GENERATED"),
            ("/ask", "services.ai_service.ask_question", {"text": "Private input", "question": "Private question"}, {"answer": "Private answer"}, "QUESTION_ANSWERED"),
            ("/podcast", "services.podcast_service.process_podcast", {"text": "Private input", "language": "twi", "host_a_voice": "Kofi", "host_b_voice": "Akua"}, ("/audio/test.wav", 1.0, [{"host": "A", "line": "Private dialogue"}]), "PODCAST_GENERATED"),
        ]
        for route, target, payload, result, expected in cases:
            with self.subTest(route=route):
                before = len(actions())
                with patch(target, side_effect=HTTPException(503, "Provider unavailable")):
                    self.assertEqual(self.client.post(route, headers=self.headers(), json=payload).status_code, 503)
                self.assertEqual(len(actions()), before)
                with patch(target, return_value=result):
                    self.assertEqual(self.client.post(route, headers=self.headers(), json=payload).status_code, 200)
                self.assertEqual(len(actions()), before + 1)
                self.assertEqual(actions()[-1], expected)
        with Session(self.engine) as db:
            for event in db.query(SystemLog).all():
                self.assertNotIn("Private", event.details or "")
                self.assertNotIn("test-password-123", event.details or "")

    def test_z_downgrade_upgrade_preserves_users_and_history(self):
        with self.engine.connect() as conn:
            users_before = conn.execute(text("SELECT count(*) FROM users")).scalar()
            conversions_before = conn.execute(text("SELECT count(*) FROM conversion_history")).scalar()
        self.migrate("downgrade", "00ab090b8bdc")
        self.migrate("upgrade", "head")
        with self.engine.connect() as conn:
            self.assertEqual(conn.execute(text("SELECT count(*) FROM users")).scalar(), users_before)
            self.assertEqual(conn.execute(text("SELECT count(*) FROM conversion_history")).scalar(), conversions_before)
            self.assertEqual(conn.execute(text("SELECT count(*) FROM conversion_history WHERE document_id IS NOT NULL")).scalar(), 0)


class PipelineRegressionTests(unittest.TestCase):
    def test_placeholder_protection_restoration_and_audio_stitching(self):
        from services import tts_service, preprocessing_service
        from pydub import AudioSegment
        wav = io.BytesIO()
        AudioSegment.silent(duration=100).export(wav, format="wav")
        sent = []

        def provider(url, json, headers):
            if url.endswith("/translate"):
                sent.append(("translate", json))
                return SimpleNamespace(status_code=200, json=lambda: json["in"])
            sent.append(("tts", json))
            return SimpleNamespace(status_code=200, content=wav.getvalue())

        with tempfile.TemporaryDirectory() as audio_dir, patch.object(settings, "AUDIO_STORAGE_PATH", audio_dir), \
                patch("services.preprocessing_service.detect_risky_words", return_value=["BERTterm"]), \
                patch("services.preprocessing_service.load_scientific_words", return_value={"Scientific"}), \
                patch("services.tts_service.requests.post", side_effect=provider):
            url, duration = tts_service.process_tts("Scientific BERTterm content.", "twi", "Kofi")
            self.assertTrue(Path(audio_dir, url.split("/")[-1]).exists())
            self.assertGreater(duration, 0)
        self.assertIn("TERMZERO", sent[0][1]["in"])
        self.assertEqual(sent[1][1]["text"], "Scientific BERTterm content.")
        self.assertEqual(sent[0][1]["lang"], "eng-twi")
        self.assertEqual(sent[1][1]["speaker_id"], "male_low")
        self.assertEqual(preprocessing_service.restore_placeholders("TERMZERO0 TERMZERO 1", {"TERMZERO0": "A", "TERMZERO1": "B"}), "A B")
        sent.clear()
        with tempfile.TemporaryDirectory() as audio_dir, patch.object(settings, "AUDIO_STORAGE_PATH", audio_dir), \
                patch("services.preprocessing_service.detect_risky_words", return_value=[]), \
                patch("services.preprocessing_service.load_scientific_words", return_value=set()), \
                patch("services.tts_service.requests.post", side_effect=provider):
            url, duration = tts_service.process_tts("Water supports life. " * 60, "twi", "Kofi")
            self.assertEqual([kind for kind, _ in sent], ["translate", "tts", "translate", "tts"])
            # WAV durations are quantized to whole audio samples.
            self.assertAlmostEqual(duration, 0.2, delta=0.001)
            self.assertTrue(Path(audio_dir, url.split("/")[-1]).exists())

    def test_supporting_ai_and_podcast_processing(self):
        from services import ai_service, podcast_service
        from pydub import AudioSegment
        result = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="Concise summary."))])
        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **kwargs: result)))
        with patch("services.ai_service.load_groq", return_value=client), \
                patch("services.preprocessing_service.detect_risky_words", return_value=[]), \
                patch("services.preprocessing_service.load_scientific_words", return_value=set()), \
                patch("services.ai_service.translate_text", side_effect=lambda text, lang: text):
            self.assertEqual(ai_service.summarise_text("Source text", "twi")["summary"], "Concise summary.")
            self.assertEqual(ai_service.ask_question("Source", "Question?")["answer"], "Concise summary.")
        wav = io.BytesIO()
        AudioSegment.silent(duration=100).export(wav, format="wav")
        result.choices[0].message.content = '[{"host":"A","line":"Welcome."},{"host":"B","line":"Discussion."}]'
        with tempfile.TemporaryDirectory() as audio_dir, patch.object(settings, "AUDIO_STORAGE_PATH", audio_dir), \
                patch("services.podcast_service.load_groq", return_value=client), \
                patch("services.preprocessing_service.detect_risky_words", return_value=[]), \
                patch("services.preprocessing_service.load_scientific_words", return_value=set()), \
                patch("services.podcast_service.translate_text", side_effect=lambda text, lang: text), \
                patch("services.podcast_service.text_to_speech", return_value=wav.getvalue()):
            url, duration, segments = podcast_service.process_podcast("Source", "twi", "Kofi", "Akua")
            self.assertTrue(Path(audio_dir, url.split("/")[-1]).exists())
            self.assertEqual(len(segments), 2)
            self.assertGreater(duration, 0)


if __name__ == "__main__":
    unittest.main()
